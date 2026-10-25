#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成外国政府动植物检疫政策变化日报(Markdown + Word + Excel)。"""
import argparse
import csv
import datetime
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates"))
from base import auto_fit_columns, auto_fit_row_heights, style_data_row, style_header_row
from normalize import DATA_DIR, REPO, biz_today, load_events, now, open_db

POLICY_TEMPLATE = os.path.join(REPO, "templates", "policy_report.md")
ACTION_ICON = {"收紧": "🔴", "放松": "🟢", "调整": "🟡", "恢复": "🔵"}
ACTION_ORDER = {"收紧": 0, "调整": 1, "放松": 2, "恢复": 3}
IMPACT_ORDER = {"高影响": 0, "中影响": 1, "低影响": 2}
POLICY_LEVEL_REVERSE = {"高影响": "high", "中影响": "medium", "低影响": "low"}
DOMAIN_LABEL = {"animal": "动物卫生", "plant": "植物保护", "both": "动植物",
                "trade": "进出口贸易", "measures": "口岸措施"}
POLICY_TABLE_HEADER = (
    "| # | 国家/地区 | 动作 | 政策领域 | 相关病害/商品 | 生效日期 | 政策状态 | 摘要 | "
    "影响类型 | 对华影响 | 后续动作 | 来源 |\n"
    "|---|---|---|---|---|---|---|---|---|---|---|---|")


def _date(value):
    """从 ISO 字符串解析日期; 失败返回 None, 不让单条脏数据拖垮整个 report。"""
    try:
        return datetime.date.fromisoformat(str(value)[:10])
    except (ValueError, TypeError):
        return None


def policy_title(policy):
    return policy.get("title_cn") or policy.get("title_en") or policy.get("summary_cn") or "（无标题）"


def policy_date(policy):
    """生效日期展示口径: 依次取 effective_date / event_date, 解析不了就明说"日期异常"。

    2026-10-19 修正(A3 闭环): 原来直接把原始字符串塞进表格与 Excel, 库里若有
    2026-99-99 / TBD 这类脏日期, 读者无法分辨那是脏数据还是真实日期。入库端已
    拒收非法日期, 这里兜住**存量脏数据**: 标成"日期异常", 不静默混进日期列。
    """
    for key in ("effective_date", "event_date"):
        raw = policy.get(key)
        if raw and _date(raw) is not None:
            return str(raw)[:10]
    return "日期异常"


def policy_impact_level(policy):
    return policy.get("impact_level") or "未研判"


def policy_dims(policy):
    """四维评分一行式: trade 4 / biosecurity 3 / ... → 3.25分; 便于追溯关注档来源。"""
    dims = policy.get("dimension_scores") or {}
    if not dims:
        return None
    parts = ["%s %s" % (k, dims.get(k, "-"))
             for k in ("trade", "biosecurity", "response", "alignment")]
    score = policy.get("impact_score")
    return "%s → %s分" % (" / ".join(parts), score if score is not None else "?")


def src_link(source):
    if not isinstance(source, dict):
        return "-"
    name, url = source.get("name") or "来源", source.get("url")
    return "[%s](%s)" % (name, url) if url else name


def sort_key(policy):
    return (ACTION_ORDER.get(policy.get("action_type"), 9),
            IMPACT_ORDER.get(policy.get("impact_level"), 3),
            str(policy.get("effective_date") or policy.get("event_date") or ""))


def collect(date, db_path=None):
    """汇总某日日报所需的记录组。

    date 必须是合法的 YYYY-MM-DD: 解析失败直接抛 ValueError(调用方转成中文错误提示
    并以非零码退出)。2026-10-19 修正: 原先 _date() 解析失败返回 None, 而本函数
    拿它做日期减法, 导致 `report.py --date 2026-13-45` / `push_report.py --date ...`
    抛裸 TypeError, 当天日报与推送一起丢(库里非空时必崩)。

    2026-10-25 移除"近期变动跟踪"桶(原 days_back=14 的 active/covered): 政策的
    后续动静由每日持续检索重新捕获, 报告不再机械重复旧条目; 历史查询走数据库
    与 webapp 面板。日报即当日快照, 所有栏目只基于今日新增(first_seen==date)。
    """
    if _date(date) is None:
        raise ValueError("日期格式非法: %r(应为 YYYY-MM-DD)" % (date,))
    con = open_db(db_path)
    all_events = [e for e in load_events(con) if e.get("record_type") == "policy"]
    con.close()
    if not all_events:
        return None
    new = [e for e in all_events if e.get("verification_status") != "merged"
           and str(e.get("first_seen", ""))[:10] == date]
    # 存量脏日期(入库端已拒收, 这里兜底): 今日新增桶里若还有解析不了的日期,
    # 保留条目(不能因为日期坏就把今天的政策从日报里抹掉), 但明确打 WARN,
    # 展示层标"日期异常"。
    for e in new:
        if policy_date(e) == "日期异常":
            print("[WARN] 今日新增记录日期无法解析(日报中显示为「日期异常」): %s @ %s "
                  "(event_date=%r, effective_date=%r)"
                  % (policy_title(e), e.get("country_en") or "?",
                     e.get("event_date"), e.get("effective_date")))
    new.sort(key=sort_key)
    return {"all": all_events, "new": new}


def policy_row(index, policy):
    domain = DOMAIN_LABEL.get(policy.get("policy_domain"), policy.get("policy_domain") or "-")
    products = "、".join(policy.get("products") or []) or \
        policy.get("disease_name_cn") or policy.get("disease_name_en") or "-"
    return "| %d | %s%s | %s%s | %s | %s | %s | %s | %s | %s | %s | %s | %s |" % (
        index, policy.get("country_cn") or "-",
        ("·" + str(policy["region"])) if policy.get("region") else "",
        ACTION_ICON.get(policy.get("action_type"), "⚪"), policy.get("action_type") or "-",
        domain, products, policy_date(policy),
        policy.get("policy_status") or "已生效", policy_title(policy)[:40],
        policy.get("impact_type") or "未研判", policy.get("impact_level") or "未研判",
        policy.get("recommended_action") or "未注明", src_link(policy.get("source")))


def make_policy_table(events):
    if not events:
        return "_（无）_\n"
    return "\n".join([POLICY_TABLE_HEADER] +
                      [policy_row(i, e) for i, e in enumerate(events, 1)]) + "\n"


def policy_headline_items(new):
    if not new:
        return ["本期无新增政策变化记录。"]
    countries = sorted({p.get("country_cn") for p in new if p.get("country_cn")})
    domains = sorted({DOMAIN_LABEL.get(p.get("policy_domain"), p.get("policy_domain"))
                      for p in new if p.get("policy_domain")})
    high = sum(1 for p in new if p.get("impact_level") == "高影响")
    medium = sum(1 for p in new if p.get("impact_level") == "中影响")
    low = sum(1 for p in new if p.get("impact_level") == "低影响")
    focus = [p for p in new if p.get("impact_level") == "高影响" or p.get("action_type") == "收紧"]
    unverified = sum(1 for p in new if p.get("verification_status") == "unverified")
    return [
        "变化了多少: 本期新增 %d 条政策变化(收紧 %d / 放松或恢复 %d / 其他 %d)。" % (
            len(new), sum(1 for p in new if p.get("action_type") == "收紧"),
            sum(1 for p in new if p.get("action_type") in ("放松", "恢复")),
            sum(1 for p in new if p.get("action_type") not in ("收紧", "放松", "恢复"))),
        "涉及国家: %d 个: %s。" % (len(countries), "、".join(countries[:8]) or "—"),
        "涉及领域: %s。" % ("、".join(domains) if domains else "—"),
        "对华影响: 高影响 %d 条 / 中影响 %d 条 / 低影响 %d 条 / 未研判 %d 条。" % (
            high, medium, low, len(new) - high - medium - low),
        "重点关注: %s" % ("; ".join("%s(%s)" % (policy_title(p)[:24], p.get("country_cn"))
                           for p in focus[:5]) if focus else "本期无。"),
    ] + (["待核实: %d 条, 见待核实信息栏。" % unverified] if unverified else [])


def make_policy_headline(new):
    return "\n".join("- " + item for item in policy_headline_items(new))


def make_policy_detail(events):
    keys = [p for p in events if p.get("action_type") in ("收紧", "调整")
            or policy_impact_level(p) == "高影响"]
    if not keys:
        return "_本期无收紧、调整或高影响类政策变化。_\n"
    blocks = []
    for i, p in enumerate(keys, 1):
        links = [src_link(p.get("source"))] + [src_link(s) for s in p.get("cross_sources", []) if isinstance(s, dict)]
        subject = "、".join(p.get("products") or []) or p.get("disease_name_cn") or p.get("disease_name_en") or "背景资料未提及"
        blocks.append(
            "### %d. %s · %s(%s)\n"
            "- **动作/状态**: %s → %s / %s\n"
            "- **摘要**: %s\n"
            "- **相关病害/商品**: %s\n"
            "- **生效日期**: %s | **有效期末**: %s\n"
            "- **影响类型/等级**: %s / %s\n"
            "- **中国关联**: %s\n"
            "%s"
            "- **建议动作**: %s\n"
            "- **依据/原文**: %s\n"
            "- **来源**: %s\n" % (
                i, p.get("country_cn") or "-", policy_title(p), p["event_id"],
                p.get("prev_action") or "此前未注明", p.get("action_type") or "-", p.get("policy_status") or "已生效",
                p.get("summary_cn") or "见来源", subject,
                policy_date(p) if (p.get("effective_date") or p.get("event_date")) else "未注明",
                p.get("effective_until") or "未注明",
                p.get("impact_type") or "未研判", p.get("impact_level") or "未研判",
                p.get("china_relevance") or "未研判",
                ("- **四维评分**: %s\n" % policy_dims(p)) if policy_dims(p) else "",
                p.get("recommended_action") or "待研判",
                p.get("legal_basis") or "见来源原文", " | ".join(links)))
    return "\n".join(blocks)


def make_policy_impact_summary(events):
    judged = [p for p in events if p.get("impact_level")]
    if not judged:
        return "_本期暂无已完成对华影响研判的政策变化。_\n"
    lines = []
    for p in sorted(judged, key=lambda x: IMPACT_ORDER.get(x.get("impact_level"), 3)):
        lines.append("- **%s(%s)** %s·%s·%s: %s" % (
            policy_title(p)[:36], p.get("country_cn"), p.get("impact_type") or "未分类",
            p.get("impact_level"), p.get("china_relevance") or "未研判",
            p.get("recommended_action") or ""))
    return "\n".join(lines) + "\n"


def make_policy_watchlist_diff(events):
    lines = []
    for p in events:
        subject = p.get("disease_name_cn") or p.get("disease_name_en") or "、".join(p.get("products") or [])
        if subject:
            lines.append("- %s(%s) %s — %s: %s" % (
                subject, p.get("country_cn"), p.get("action_type") or "-",
                p.get("impact_level") or "未研判", p.get("summary_cn") or policy_title(p)[:40]))
    return "\n".join(lines) + "\n" if lines else "_本期无记录明确关联病害或商品的政策变化。_\n"


def make_open_items(events):
    icons = {"unverified": "❓", "false_positive": "⚠️"}
    lines = []
    for e in events:
        status, name = e.get("verification_status"), policy_title(e)
        if status in icons:
            lines.append("- %s %s @ %s(%s) 核验:%s: %s" % (
                icons[status], name, e.get("country_cn"), e["event_id"], status, e.get("summary_cn") or "见来源"))
        elif status == "verified" and not e.get("impact_level"):
            lines.append("- ⏳ %s @ %s(%s) 已核验、待影响研判: %s" % (
                name, e.get("country_cn"), e["event_id"], e.get("summary_cn") or "见来源"))
    return "\n".join(lines) + "\n" if lines else "_无待核实/待研判政策记录。_\n"


def make_stats(all_events, new):
    by_status = {}
    for e in all_events:
        by_status[e.get("verification_status")] = by_status.get(e.get("verification_status"), 0) + 1
    detail = " / ".join("%s %d" % (k, v) for k, v in sorted(by_status.items(), key=lambda kv: -kv[1]))
    return "- 政策库累计: %d 条(%s)\n- 本期新增: %d 条\n" % (
        len(all_events), detail, len(new))


def make_sources(events):
    seen, items = {}, []
    for e in events:
        for s in [e.get("source")] + list(e.get("cross_sources") or []):
            if isinstance(s, dict) and s.get("url") and s["url"] not in seen:
                seen[s["url"]] = e["event_id"]
                items.append("- [%s](%s) — 政策 %s" % (s.get("name") or s["url"], s["url"], e["event_id"]))
    return "\n".join(items) + "\n" if items else "_（本期无来源记录）_\n"


def export_table(base_path, events):
    def val(value): return "未注明" if value in (None, "", [], {}) else value
    headers = ["记录ID", "政策标题", "英文标题", "动作类型", "政策领域", "发布国家/地区", "发布机构",
               "涉及国家/地区", "受影响商品", "关联病害", "适用范围", "发布日期", "生效日期", "有效期",
               "政策状态", "公告/法规编号", "核验状态", "影响类型", "影响等级", "关注档", "四维评分",
               "中国关联", "建议动作",
               "原文摘要", "来源名称", "来源URL", "来源层级", "抓取时间"]
    rows = [[val(e.get("event_id")), val(e.get("title_cn")), val(e.get("title_en")),
             val(e.get("action_type")), val(e.get("policy_domain")), val(e.get("country_cn")),
             val(e.get("issuer_cn") or e.get("issuer_en")), "、".join(e.get("target_countries") or []) or val(None),
             "、".join(e.get("products") or []) or val(None), val(e.get("disease_name_cn") or e.get("disease_name_en")),
             val(e.get("scope")), val(e.get("published_date") or e.get("report_date") or (e.get("source") or {}).get("publish_date")),
             val(policy_date(e)), val(e.get("effective_until")), val(e.get("policy_status") or "已生效"),
             val(e.get("legal_basis")), val(e.get("verification_status")), val(e.get("impact_type")), val(e.get("impact_level")),
             val(e.get("impact_focus")), val(policy_dims(e)),
             val(e.get("china_relevance")), val(e.get("recommended_action")), val(e.get("summary_cn")),
             val((e.get("source") or {}).get("name")), val((e.get("source") or {}).get("url")), val((e.get("source") or {}).get("tier")),
             val(e.get("first_seen"))] for e in events]
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook(); wb.properties.title = "全球动植物检疫政策监测台账"; wb.properties.creator = "疫见全球"
        ws = wb.active; ws.title = "政策变化台账"; ws.freeze_panes = "A2"
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"; ws.append(headers)
        for row in rows: ws.append(row)
        style_header_row(ws, 1, 1, len(headers))
        for idx in range(len(rows)): style_data_row(ws, idx + 2, 1, len(headers), idx)
        auto_fit_columns(ws, min_width=10, max_width=34); auto_fit_row_heights(ws, 1, 2)
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)
        if rows: ws.conditional_formatting.add("A2:%s%d" % (get_column_letter(len(headers)), len(rows) + 1), FormulaRule(formula=["MOD(ROW(),2)=0"], fill=PatternFill("solid", fgColor="F7F7F5")))
        path = base_path + ".xlsx"; wb.save(path); return path
    except ImportError:
        path = base_path + ".csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as f: csv.writer(f).writerows([headers] + rows)
        return path


def ensure_docx(date, db_path=None):
    """Word 日报缺失时补生成一次; 已存在则原样保留(幂等)。返回 0=可用, 1=仍不可用。

    为什么要有这个入口(2026-10-20 生产现场发现):
      Word 日报由 `report.py --docx` 生成, 而那一步是 **Agent 照 prompt 执行** 的
      —— 实测 10-20 的跑批漏掉了 --docx, 只产出 csv/md/xlsx。而推送侧当时是
      `if isfile(docx_path)` **静默跳过附件**: 用户收到的是纯文本, 以为带 Word,
      事后从推送记录里完全看不出来(钉钉/微信都收过"缺附件"的日报)。

      现在补两道确定性兜底 —— 跑批结束(run_scan.sh)与推送前(push_report.py)
      各调本入口一次: 缺就补, 补不上就明确告警, 不再假装有附件。
    """
    path = os.path.join(DATA_DIR, "reports", "%s-policy-report.docx" % date)
    if os.path.isfile(path) and os.path.getsize(path) > 0:
        print("[OK] Word 日报已存在: %s" % path)
        return 0
    # 同源的 Markdown 都没有 → 整份日报根本没生成, 补一个 Word 只会造出"空壳日报"
    # 掩盖真实失败(10-22 实测: 检索全废、一条没入库, 却兜底生成了 62KB 的 Word,
    # 看上去像"报告已出", 实际当天 0 条记录 —— 比缺文件更坏)。这种情况报失败。
    md_path = os.path.join(DATA_DIR, "reports", "%s-policy-report.md" % date)
    if not (os.path.isfile(md_path) and os.path.getsize(md_path) > 0):
        print("[WARN] 当日日报未生成(缺 %s), 不补 Word 日报 —— "
              "补出来会是空壳, 掩盖真实失败。" % os.path.basename(md_path))
        return 1
    print("[提示] Word 日报缺失, 正在补生成: %s" % path)
    try:
        import report_docx
        # 显式给 out_dir: 与上面的检查路径保持同一个目录, 避免 report_docx
        # 用它自己导入的 DATA_DIR 写到别处去(生产上 DATA_DIR 可被环境变量改写)。
        print("[OK] %s" % report_docx.build(date, out_dir=os.path.join(DATA_DIR, "reports"),
                                           db_path=db_path))
    except ImportError:
        print("[WARN] 未安装 python-docx, 无法生成 Word 简报(推送将只带正文)")
        return 1
    except SystemExit as exc:
        # report_docx 在"政策库为空"时抛的是 SystemExit(不是 Exception) ——
        # 不接住的话会直接中断调用方(推送流程被莫名打断, 退出码还很难解释)。
        print("[WARN] Word 简报生成中止: %s" % (exc or "未知原因"))
        return 1
    except Exception as exc:
        print("[WARN] Word 简报生成失败: %s" % exc)
        return 1
    if os.path.isfile(path) and os.path.getsize(path) > 0:
        return 0
    print("[WARN] Word 简报生成后仍不可用: %s" % path)
    return 1


def main():
    ap = argparse.ArgumentParser(description="生成政策监测日报")
    ap.add_argument("--date", default=biz_today())
    ap.add_argument("--policy", action="store_true", help="(已废弃, 无效果) 政策为唯一模式, 兼容旧定时/面板命令")
    ap.add_argument("--excel", action="store_true"); ap.add_argument("--docx", action="store_true"); ap.add_argument("--db-path")
    ap.add_argument("--ensure-docx", action="store_true",
                    help="Word 日报缺失时才补生成一次(跑批后/推送前的兜底; 已有则不动)")
    args = ap.parse_args()
    if args.ensure_docx:
        return ensure_docx(args.date, args.db_path)
    try:
        data = collect(args.date, args.db_path)
    except ValueError as exc:
        print("[错误] %s" % exc)
        return 2
    if data is None: data = {"all": [], "new": []}
    new = data["new"]
    replacements = {"{{date}}": args.date, "{{generated_at}}": now(), "{{headline}}": make_policy_headline(new), "{{new_table}}": make_policy_table(new), "{{focus_detail}}": make_policy_detail(new), "{{impact_summary}}": make_policy_impact_summary(new), "{{watchlist_diff}}": make_policy_watchlist_diff(new), "{{unverified}}": make_open_items(new), "{{stats}}": make_stats(data["all"], new), "{{sources}}": make_sources(new)}
    base = os.path.join(DATA_DIR, "reports", "%s-policy-report" % args.date)
    with open(POLICY_TEMPLATE, encoding="utf-8") as f: report_md = f.read()
    for key, value in replacements.items(): report_md = report_md.replace(key, value)
    os.makedirs(os.path.dirname(base), exist_ok=True)
    with open(base + ".md", "w", encoding="utf-8") as f: f.write(report_md)
    print("[OK] %s.md" % base)
    if args.excel: print("[OK] %s" % export_table(base, new))
    if args.docx:
        try:
            import report_docx; print("[OK] %s" % report_docx.build(args.date, data))
        except ImportError: print("[WARN] 未安装 python-docx, 跳过 Word 简报")
        except SystemExit as exc: print("[WARN] Word 简报生成中止: %s" % (exc or "未知原因"))
        except Exception as exc: print("[WARN] Word 简报生成失败: %s" % exc)
    print("\n本期: 新增 %d 条(政策库共 %d 条)。" % (len(new), len(data["all"])))
    return 0


if __name__ == "__main__": raise SystemExit(main())
