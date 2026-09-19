#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成外国政府动植物检疫政策监测 Word 简报(.docx)。

用法:
  python report.py --date 2026-09-12 --docx       # 默认政策监测版
  python report_docx.py --date 2026-09-12
输出: data/reports/<date>-policy-report.docx(默认)
"""
import argparse
import datetime
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import report
from normalize import DATA_DIR

RISK_FILL = {"high": "FDE9E7", "medium": "FFF4DE", "low": "E6F4EA"}
RISK_TEXT = {"high": "高", "medium": "中", "low": "低"}
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xC0, 0x39, 0x2B)


def _run(p, text, size=10.5, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return r


def _para(doc, text="", size=10.5, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        _run(p, text, size=size, bold=bold, color=color)
    return p


def _heading(doc, text):
    return _para(doc, text, size=13, bold=True, color=ACCENT, space_after=6)


def _bullet(doc, label, text=None, color=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if text is None:
        _run(p, label, color=color)
    else:
        if label:
            _run(p, label, bold=True)
        _run(p, text, color=color)
    return p


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_text(cell, text, bold=False):
    cell.text = ""
    _run(cell.paragraphs[0], str(text), size=9.5, bold=bold)


def add_hyperlink(paragraph, url, text, size=9.5):
    """python-docx 无内置超链接, 手工拼 w:hyperlink。"""
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    for el in (fonts, color, u, sz):
        r_pr.append(el)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)



def _policy_table(doc, events):
    cols = ["政策标题", "动作", "领域", "国家/地区", "商品/病害", "生效日期", "政策状态",
            "影响类型", "影响等级", "后续动作"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        _cell_text(t.rows[0].cells[i], c, bold=True)
    for e in events:
        title = e.get("title_cn") or e.get("title_en") or e.get("summary_cn") or "（无标题）"
        domain = e.get("policy_domain") or "-"
        subject = "、".join(e.get("products") or []) or e.get("disease_name_cn") or e.get("disease_name_en") or "-"
        level = report.policy_impact_level(e) or "未研判"
        vals = [title, e.get("action_type") or "-", domain,
                "%s%s" % (e.get("country_cn") or "-", ("·" + str(e["region"])) if e.get("region") else ""),
                subject, e.get("effective_date") or e.get("event_date") or "-",
                e.get("policy_status") or "已生效", e.get("impact_type") or "未研判",
                level, e.get("recommended_action") or "未注明"]
        row = t.add_row().cells
        for i, v in enumerate(vals):
            _cell_text(row[i], v)
        fill = RISK_FILL.get(report.POLICY_LEVEL_REVERSE.get(level))
        if fill:
            _shade(row[8], fill)
    return t


def _build_policy(date, data=None, out_dir=None, db_path=None):
    """生成政策监测 Word 简报。"""
    if data is None:
        data = report.collect(date, 14, db_path)
    if data is None:
        raise SystemExit("[提示] 政策库为空: 先运行 normalize.py 入库政策记录。")
    new, active, covered = data["new"], data["active"], data["covered"]
    doc = Document()
    _para(doc, "全球动植物检疫政策监测日报", size=18, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, "疫见全球 · Government Phytosanitary & Animal Health Policy Monitor   |   %s   |   生成于 %s" % (
        date, datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    sec = 0
    def _h(title):
        nonlocal sec
        sec += 1
        _heading(doc, "%s、%s" % ("一二三四五六七八九十"[sec - 1], title))

    _h("本期政策变化速览")
    for item in report.policy_headline_items(new, active):
        _bullet(doc, item)
    _h("本期新增政策变化(%d 条)" % len(new))
    _policy_table(doc, new)
    if active:
        _h("近期变动跟踪(%d 条)" % len(active))
        _policy_table(doc, active)

    _h("收紧与调整动作详情")
    focus = [e for e in covered if e.get("action_type") in ("收紧", "调整")
             or e.get("impact_level") == "高影响"]
    if not focus:
        _para(doc, "本期无收紧、调整或高影响类政策变化。", color=GRAY)
    for i, e in enumerate(focus, 1):
        title = e.get("title_cn") or e.get("title_en") or "（无标题）"
        _para(doc, "%d. %s · %s(%s)" % (i, e.get("country_cn") or "-", title, e["event_id"]),
              bold=True, size=11, space_after=2)
        _bullet(doc, "动作: ", "%s → %s" % (e.get("prev_action") or "此前未注明", e.get("action_type") or "-"))
        _bullet(doc, "摘要: ", e.get("summary_cn") or "见来源")
        _bullet(doc, "领域/对象: ", "%s / %s" % (
            e.get("policy_domain") or "-", "、".join(e.get("products") or []) or
            e.get("disease_name_cn") or e.get("disease_name_en") or "未注明"))
        _bullet(doc, "发布机构/法律依据: ", "%s / %s" % (
            e.get("issuer_cn") or e.get("issuer_en") or "未注明", e.get("legal_basis") or "未注明"))
        _bullet(doc, "生效/状态: ", "%s / %s" % (
            e.get("effective_date") or e.get("event_date") or "未注明", e.get("policy_status") or "已生效"))
        _bullet(doc, "对华影响: ", "%s — %s" % (
            e.get("impact_level") or "未研判", e.get("impact_rationale") or "待研判"))
        _bullet(doc, "影响类型/中国关联: ", "%s / %s" % (
            e.get("impact_type") or "未研判", e.get("china_relevance") or "未研判"))
        _bullet(doc, "建议动作: ", e.get("recommended_action") or "待研判")
        pnode = doc.add_paragraph(style="List Bullet")
        pnode.paragraph_format.space_after = Pt(3)
        _run(pnode, "来源: ")
        for s in [e.get("source")] + list(e.get("cross_sources") or []):
            if isinstance(s, dict) and s.get("url"):
                add_hyperlink(pnode, s["url"], s.get("name") or s["url"])
                _run(pnode, "   ")

    _h("对华影响研判综述")
    judged = [e for e in covered if e.get("impact_level")]
    if not judged:
        _para(doc, "本期暂无已完成对华影响研判的政策变化。", color=GRAY)
    for e in judged:
        _bullet(doc, "%s(%s): " % (e.get("title_cn") or e.get("title_en") or "（无标题）", e.get("country_cn") or "-"),
                "%s·%s·%s — %s" % (e.get("impact_type") or "未分类", e.get("impact_level"),
                                   e.get("china_relevance") or "未研判",
                                   e.get("impact_rationale") or e.get("recommended_action") or ""))

    _h("待核实信息")
    open_rows = [e for e in covered if e.get("verification_status") in ("unverified", "false_positive")
                 or (e.get("verification_status") == "verified" and not e.get("impact_level"))]
    if not open_rows:
        _para(doc, "无待核实/待研判政策记录。", color=GRAY)
    for e in open_rows:
        st = e.get("verification_status")
        tag = {"unverified": "未核验", "false_positive": "存疑"}.get(st, "待研判")
        _bullet(doc, "[%s] " % tag,
                "%s @ %s(%s): %s" % (e.get("title_cn") or e.get("title_en") or "（无标题）",
                                     e.get("country_cn") or "-", e["event_id"], e.get("summary_cn") or "见来源"),
                color=RED if st == "false_positive" else None)

    _h("来源索引")
    seen, idx = set(), 0
    for e in covered:
        for s in [e.get("source")] + list(e.get("cross_sources") or []):
            if isinstance(s, dict) and s.get("url") and s["url"] not in seen:
                seen.add(s["url"]); idx += 1
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                _run(p, "%d. " % idx, size=9.5)
                add_hyperlink(p, s["url"], "%s — 政策 %s" % (s.get("name") or s["url"], e["event_id"]))
    _para(doc, "", space_after=8)
    _para(doc, "声明: 本简报由 AI 辅助生成,所有政策变化附原始来源;核验状态与对华影响仅供情报参考,不构成决策或执法依据。中国海关总署等官方公告以正式发布为准。", size=8.5, color=GRAY)
    out_dir = out_dir or os.path.join(DATA_DIR, "reports")
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "%s-policy-report.docx" % date)
    doc.save(path)
    return path


def build(date, data=None, out_dir=None, db_path=None):
    return _build_policy(date, data=data, out_dir=out_dir, db_path=db_path)


def main():
    ap = argparse.ArgumentParser(description="生成 Word 情报简报(数据与 Markdown 日报同源)")
    ap.add_argument("--date", default=datetime.date.today().isoformat())
    ap.add_argument("--days-back", type=int, default=14)
    ap.add_argument("--db-path", help="SQLite 路径覆盖")
    ap.add_argument("--out-dir", help="输出目录(默认 data/reports)")
    args = ap.parse_args()
    path = build(args.date, out_dir=args.out_dir, db_path=args.db_path)
    print("[OK] %s" % path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
