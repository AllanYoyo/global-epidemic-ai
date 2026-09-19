"""轻量 Excel 台账样式工具,供政策/疫情导出共用。"""
import platform
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PRIMARY = "1B2A4A"
PRIMARY_LIGHT = "D6E4F0"
NEUTRAL_900 = "37352F"
NEUTRAL_200 = "E9E9E8"
NEUTRAL_100 = "F7F7F5"
NEUTRAL_0 = "FFFFFF"
FONT_NAME = "Microsoft YaHei" if platform.system() == "Windows" else "Noto Sans CJK SC"
HEADER_BOLD = True


def setup_sheet(ws, title, last_col):
    ws.title = title
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = "A1:%s1048576" % _col(last_col)


def style_header_row(ws, row_num=1, col_start=1, col_end=None):
    col_end = col_end or ws.max_column
    fill = PatternFill("solid", fgColor=PRIMARY)
    font = Font(name=FONT_NAME, size=10, bold=HEADER_BOLD, color="FFFFFF")
    border = Border(bottom=Side(style="thin", color=NEUTRAL_200))
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row_num, col)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row_num].height = 30


def style_data_row(ws, row_num, col_start=1, col_end=None, row_index=0):
    col_end = col_end or ws.max_column
    fill = NEUTRAL_0 if row_index % 2 == 0 else NEUTRAL_100
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row_num, col)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def auto_fit_columns(ws, min_width=10, max_width=34):
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, col).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
        width = max([len(v) for v in values] or [0]) * 1.15 + 2
        ws.column_dimensions[_col(col)].width = min(max(width, min_width), max_width)


def auto_fit_row_heights(ws, header_row=1, data_start_row=2):
    for row in range(data_start_row, ws.max_row + 1):
        max_lines = 1
        for cell in ws[row]:
            value = str(cell.value or "")
            max_lines = max(max_lines, value.count("\n") + 1, (len(value) // 32) + 1)
        ws.row_dimensions[row].height = min(max(22, 15 * max_lines), 90)


def _col(index):
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result
